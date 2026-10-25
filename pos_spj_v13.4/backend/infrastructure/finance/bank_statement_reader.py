"""Lector de estados de cuenta CSV/XLSX (§15).

Convierte un archivo del banco en partidas ``{"transaction_date", "description",
"amount", "external_reference"}`` con importe CON SIGNO (positivo entra,
negativo sale). Reconoce encabezados en español o inglés; acepta una columna de
importe con signo o dos columnas cargo/abono. Los importes se leen como
Decimal desde texto (nunca float). Un renglón ilegible detiene la importación
con su número de renglón: no se importa a medias.

La interfaz para API bancaria futura es el mismo resultado: una lista de
partidas con el mismo contrato.
"""

from __future__ import annotations

import csv
import io
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from backend.domain.finance.exceptions import FinanceDomainError

_DATE = ("fecha", "date", "fecha operacion", "fecha de operacion", "transaction date")
_DESCRIPTION = ("descripcion", "concepto", "description", "detalle", "movimiento")
_AMOUNT = ("importe", "monto", "amount", "cantidad")
_DEBIT = ("cargo", "cargos", "retiro", "retiros", "debit", "withdrawal")
_CREDIT = ("abono", "abonos", "deposito", "depositos", "credit", "deposit")
_REFERENCE = ("referencia", "reference", "folio", "rastreo", "clave de rastreo")
_DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%Y/%m/%d")


@dataclass(frozen=True, slots=True)
class ParsedStatement:
    lines: list[dict]
    source_format: str
    file_name: str


def _normalize(text: str) -> str:
    folded = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return " ".join(folded.lower().replace("_", " ").split())


def _find(headers: list[str], options: tuple[str, ...]) -> int | None:
    for index, header in enumerate(headers):
        if _normalize(header) in options:
            return index
    return None


def _decimal(raw, row_number: int) -> Decimal | None:
    if raw is None:
        return None
    if isinstance(raw, float):
        raw = repr(raw)
    text = str(raw).strip().replace("$", "").replace(",", "").replace(" ", "")
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    try:
        value = Decimal(text)
    except InvalidOperation as exc:
        raise FinanceDomainError(f"Renglón {row_number}: importe ilegible {raw!r}") from exc
    return -value if negative else value


def _date(raw, row_number: int) -> date:
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw or "").strip()[:10]
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise FinanceDomainError(f"Renglón {row_number}: fecha ilegible {raw!r}")


def _rows_to_lines(rows: list[list], first_data_row: int) -> list[dict]:
    if not rows:
        raise FinanceDomainError("El archivo no tiene renglones")
    headers = [str(cell or "") for cell in rows[0]]
    date_col = _find(headers, _DATE)
    amount_col = _find(headers, _AMOUNT)
    debit_col, credit_col = _find(headers, _DEBIT), _find(headers, _CREDIT)
    description_col = _find(headers, _DESCRIPTION)
    reference_col = _find(headers, _REFERENCE)
    if date_col is None or (amount_col is None and (debit_col is None or credit_col is None)):
        raise FinanceDomainError(
            "Encabezados no reconocidos: se requiere Fecha y un Importe con signo, o "
            "columnas Cargo y Abono")
    lines: list[dict] = []
    for offset, row in enumerate(rows[1:]):
        number = first_data_row + offset
        if not any(str(cell or "").strip() for cell in row):
            continue

        def cell(index):
            return row[index] if index is not None and index < len(row) else None

        if amount_col is not None:
            amount = _decimal(cell(amount_col), number)
        else:
            debit = _decimal(cell(debit_col), number) or Decimal("0")
            credit = _decimal(cell(credit_col), number) or Decimal("0")
            amount = credit - abs(debit)
        if amount is None or amount == 0:
            raise FinanceDomainError(f"Renglón {number}: importe vacío o cero")
        lines.append({
            "transaction_date": _date(cell(date_col), number).isoformat(),
            "description": str(cell(description_col) or "").strip(),
            "amount": format(amount, "f"),
            "external_reference": str(cell(reference_col) or "").strip(),
        })
    if not lines:
        raise FinanceDomainError("El archivo no tiene partidas")
    return lines


def read_statement_file(path: str | Path) -> ParsedStatement:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix == ".csv":
        raw = file_path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t|")
        rows = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
        return ParsedStatement(_rows_to_lines(rows, 2), "CSV", file_path.name)
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
        return ParsedStatement(_rows_to_lines(rows, 2), "XLSX", file_path.name)
    raise FinanceDomainError("Formato no soportado: use CSV o XLSX")


def parse_manual_lines(text: str) -> list[dict]:
    """Captura controlada: ``AAAA-MM-DD;descripción;importe;referencia`` por renglón."""
    rows = [["fecha", "descripcion", "importe", "referencia"]]
    for raw in (text or "").splitlines():
        if raw.strip():
            rows.append([part.strip() for part in raw.split(";")])
    return _rows_to_lines(rows, 1)
